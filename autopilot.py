#!/usr/bin/env python3
"""questionnaire-autopilot: draft answers to security/vendor questionnaires from
your own policy docs via deterministic TF-IDF retrieval + an LLM, with citations,
confidence, and a needs-review flag.

Every output is labelled an AI-generated draft requiring human review.
"""
import argparse
import csv
import glob
import os
import re
from pathlib import Path

from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import linear_kernel

import llm

AI_DRAFT_HEADER = (
    "AI-GENERATED DRAFT ANSWERS - these are a drafting aid produced from your "
    "policy documents and MUST be reviewed, corrected, and approved by a "
    "qualified person before use. Not a compliance guarantee."
)

COLUMNS = ["Question", "Draft Answer", "Cited Source", "Confidence", "Needs Review"]

# --- Calibration knobs (TF-IDF cosine scores run low for short docs) ---------
# ponytail: naive linear calibration; swap for score normalisation if a wider
# policy corpus skews the distribution.
NO_MATCH_THRESHOLD = 0.05   # top cosine below this = "no relevant policy found"
FULL_CONFIDENCE_AT = 0.45   # cosine at/above this maps to confidence 1.0
LOW_CONFIDENCE_CUTOFF = 0.5  # confidence below this -> needs_review
HEDGE_PENALTY = 0.5         # multiply confidence when the model hedged

HEDGE_RE = re.compile(
    r"\b(may|might|possibly|unclear|cannot|unable|do(?:es)? not (?:appear|seem|"
    r"clearly|address|cover|specify)|not (?:specified|stated|mentioned|found|"
    r"clear|address)|no relevant|insufficient|unknown|unsure)\b",
    re.I,
)


# --- Loading & chunking ------------------------------------------------------
def chunk_policy(text, source):
    """Split a markdown policy into (source, section, body) chunks by heading."""
    chunks, section, buf = [], source, []

    def flush():
        body = "\n".join(buf).strip()
        if body:
            chunks.append((source, section, body))

    for line in text.splitlines():
        if line.lstrip().startswith("#"):
            flush()
            section = line.lstrip("#").strip() or source
            buf = []
        else:
            buf.append(line)
    flush()
    return chunks


def load_policies(policies_dir):
    """Return a list of chunk dicts {source, section, text} from *.md files."""
    chunks = []
    for path in sorted(glob.glob(os.path.join(policies_dir, "*.md"))):
        text = Path(path).read_text(encoding="utf-8")
        name = Path(path).name
        for _, section, body in chunk_policy(text, Path(path).stem):
            chunks.append({"source": name, "section": section, "text": body})
    return chunks


# --- Retrieval (deterministic TF-IDF) ----------------------------------------
class Retriever:
    def __init__(self, chunks):
        if not chunks:
            raise ValueError("No policy chunks to index.")
        self.chunks = chunks
        corpus = [f"{c['section']}. {c['text']}" for c in chunks]
        self.vectorizer = TfidfVectorizer(stop_words="english")
        self.matrix = self.vectorizer.fit_transform(corpus)  # l2-normalised

    def retrieve(self, question, k=3):
        """Return [(chunk, cosine_score), ...] for the top-k chunks."""
        q = self.vectorizer.transform([question])
        sims = linear_kernel(q, self.matrix).ravel()  # cosine (both l2-normalised)
        order = sims.argsort()[::-1][:k]
        return [(self.chunks[i], float(sims[i])) for i in order]


# --- Scoring -----------------------------------------------------------------
def score_confidence(top_score, hedged, has_context):
    """Confidence in [0,1] from retrieval strength, penalised if the model hedged."""
    if not has_context:
        return 0.0
    base = min(top_score / FULL_CONFIDENCE_AT, 1.0)
    if hedged:
        base *= HEDGE_PENALTY
    return round(base, 2)


# --- Drafting ----------------------------------------------------------------
def _build_messages(question, retrieved):
    context = "\n\n".join(
        f"[{c['source']} > {c['section']}]\n{c['text']}"
        for c, s in retrieved
        if s >= NO_MATCH_THRESHOLD
    )
    system = (
        "You are a security compliance analyst drafting answers to a vendor "
        "security questionnaire. Use ONLY the company policy excerpts provided. "
        "Write a concise, factual answer in 2-5 sentences. If the excerpts do not "
        "clearly address the question, reply exactly that 'The provided policies do "
        "not clearly address this' and note what is missing. Never invent controls."
    )
    user = f"Question: {question}\n\nPolicy excerpts:\n{context}\n\nDraft answer:"
    return [{"role": "system", "content": system}, {"role": "user", "content": user}]


def draft_answer(question, retrieved, chat_fn=llm.chat):
    """Draft one answer dict for a question given its retrieved chunks."""
    top = retrieved[0] if retrieved else None
    top_score = top[1] if top else 0.0
    has_context = top_score >= NO_MATCH_THRESHOLD

    if not has_context:
        return {
            "Question": question,
            "Draft Answer": "No relevant policy content was found. A qualified "
            "reviewer must answer this manually.",
            "Cited Source": "- none found -",
            "Confidence": 0.0,
            "Needs Review": True,
        }

    try:
        answer = chat_fn(_build_messages(question, retrieved)).strip()
    except llm.LLMError as e:
        answer = (
            "[LLM UNAVAILABLE - review the cited policy excerpt and draft manually]\n"
            f"{top[0]['text']}\n\n(reason: {e})"
        )
        hedged = True
    else:
        hedged = bool(HEDGE_RE.search(answer))

    confidence = score_confidence(top_score, hedged, has_context)
    return {
        "Question": question,
        "Draft Answer": answer,
        "Cited Source": f"{top[0]['source']} > {top[0]['section']}",
        "Confidence": confidence,
        # Hedging (or an unavailable model) forces review even when retrieval was strong.
        "Needs Review": hedged or confidence < LOW_CONFIDENCE_CUTOFF,
    }


# --- Input / output ----------------------------------------------------------
def read_questions(path):
    """Read questions from a .csv or .xlsx file with a 'Question' column."""
    ext = Path(path).suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        rows = list(wb.active.iter_rows(values_only=True))
        header = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
        qi = header.index("question")
        return [str(r[qi]).strip() for r in rows[1:] if r and r[qi]]

    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        key = next((c for c in reader.fieldnames or [] if c.strip().lower() == "question"), None)
        if key is None:
            raise ValueError(f"No 'Question' column found in {path}")
        return [row[key].strip() for row in reader if row.get(key, "").strip()]


def write_csv(path, results):
    with open(path, "w", newline="", encoding="utf-8") as f:
        f.write(f"# {AI_DRAFT_HEADER}\n")
        writer = csv.DictWriter(f, fieldnames=COLUMNS)
        writer.writeheader()
        for r in results:
            writer.writerow({c: r[c] for c in COLUMNS})


def write_xlsx(path, results):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Answers"
    warn = Font(bold=True, color="9C0006")
    ws.append([AI_DRAFT_HEADER])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))
    ws.cell(1, 1).font = warn
    ws.append(COLUMNS)
    for c in ws[2]:
        c.font = Font(bold=True)
    flag_fill = PatternFill("solid", fgColor="FFF2CC")
    for r in results:
        ws.append([r[c] for c in COLUMNS])
        if r["Needs Review"]:
            for c in ws[ws.max_row]:
                c.fill = flag_fill
    widths = {"A": 50, "B": 70, "C": 32, "D": 12, "E": 13}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=3):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def write_report(path, results):
    total = len(results)
    high = sum(1 for r in results if r["Confidence"] >= 0.7)
    low = sum(1 for r in results if r["Confidence"] < 0.4)
    medium = total - high - low
    flagged = [r for r in results if r["Needs Review"]]
    lines = [
        f"> **{AI_DRAFT_HEADER}**",
        "",
        "# Questionnaire Autopilot - Review Report",
        "",
        f"- Total questions: **{total}**",
        f"- High confidence (>= 0.70): **{high}**",
        f"- Medium confidence (0.40-0.69): **{medium}**",
        f"- Low confidence (< 0.40): **{low}**",
        f"- Flagged **needs review**: **{len(flagged)}**",
        "",
        "## Questions flagged for review",
        "",
    ]
    if flagged:
        for r in flagged:
            lines.append(f"- ({r['Confidence']:.2f}) {r['Question']}  -  source: {r['Cited Source']}")
    else:
        lines.append("- None. Still have a qualified person spot-check every answer.")
    lines.append("")
    Path(path).write_text("\n".join(lines), encoding="utf-8")


# --- Pipeline ----------------------------------------------------------------
def run(questions_path, policies_dir, out_dir, k=3, chat_fn=llm.chat):
    questions = read_questions(questions_path)
    retriever = Retriever(load_policies(policies_dir))
    results = [draft_answer(q, retriever.retrieve(q, k), chat_fn) for q in questions]

    os.makedirs(out_dir, exist_ok=True)
    write_csv(os.path.join(out_dir, "answers.csv"), results)
    write_xlsx(os.path.join(out_dir, "answers.xlsx"), results)
    write_report(os.path.join(out_dir, "review-report.md"), results)
    return results


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--questions", default="input/questions.csv")
    ap.add_argument("--policies", default="policies")
    ap.add_argument("--out", default="output")
    ap.add_argument("-k", type=int, default=3, help="chunks retrieved per question")
    args = ap.parse_args()

    results = run(args.questions, args.policies, args.out, k=args.k)
    flagged = sum(1 for r in results if r["Needs Review"])
    print(f"Drafted {len(results)} answers -> {args.out}/ "
          f"({flagged} flagged for review). Wrote answers.xlsx, answers.csv, review-report.md.")
    print("Reminder: AI-generated drafts. A qualified person must review before use.")


if __name__ == "__main__":
    main()
