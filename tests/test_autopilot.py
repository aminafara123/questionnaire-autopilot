"""Offline tests for the deterministic core. No network, no API key, LLM mocked."""
import csv
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import autopilot as ap

POLICIES = str(Path(__file__).resolve().parents[1] / "policies")


def fake_chat(answer="MFA is enforced for all production access using TOTP or FIDO2 keys."):
    return lambda messages: answer


# --- Retrieval ---------------------------------------------------------------
def test_retrieval_ranks_relevant_policy_first():
    r = ap.Retriever(ap.load_policies(POLICIES))
    top_chunk, score = r.retrieve("Is data encrypted at rest and in transit?", k=3)[0]
    assert "encryption" in top_chunk["source"].lower()
    assert score > ap.NO_MATCH_THRESHOLD


def test_retrieval_is_deterministic():
    r = ap.Retriever(ap.load_policies(POLICIES))
    q = "How do you handle disaster recovery and backups?"
    assert r.retrieve(q)[0][0]["source"] == r.retrieve(q)[0][0]["source"]


def test_chunking_splits_by_heading():
    chunks = ap.chunk_policy("# Title\nintro\n## Section A\nbody a\n## Section B\nbody b", "doc")
    sections = {c[1] for c in chunks}
    assert {"Section A", "Section B"} <= sections


# --- Confidence scoring ------------------------------------------------------
def test_confidence_zero_without_context():
    assert ap.score_confidence(0.9, hedged=False, has_context=False) == 0.0


def test_confidence_rises_with_retrieval_score():
    assert ap.score_confidence(0.1, False, True) < ap.score_confidence(0.4, False, True)


def test_hedging_lowers_confidence():
    assert ap.score_confidence(0.4, True, True) < ap.score_confidence(0.4, False, True)


# --- needs_review logic ------------------------------------------------------
def test_no_matching_policy_flags_needs_review():
    r = ap.Retriever(ap.load_policies(POLICIES))
    q = "What is your approach to interplanetary spacecraft propulsion telemetry?"
    result = ap.draft_answer(q, r.retrieve(q), chat_fn=fake_chat())
    assert result["Needs Review"] is True
    assert result["Confidence"] == 0.0
    assert result["Cited Source"] == "- none found -"


def test_strong_match_not_flagged():
    r = ap.Retriever(ap.load_policies(POLICIES))
    q = "Do you enforce multi-factor authentication for production systems?"
    result = ap.draft_answer(q, r.retrieve(q), chat_fn=fake_chat())
    assert result["Needs Review"] is False
    assert "access-control" in result["Cited Source"]


def test_hedged_answer_flags_review():
    r = ap.Retriever(ap.load_policies(POLICIES))
    q = "Do you enforce multi-factor authentication for production systems?"
    hedged = fake_chat("The provided policies do not clearly address this; MFA is not specified.")
    result = ap.draft_answer(q, r.retrieve(q), chat_fn=hedged)
    assert result["Needs Review"] is True


def test_llm_unavailable_degrades_gracefully():
    r = ap.Retriever(ap.load_policies(POLICIES))
    q = "Do you enforce multi-factor authentication?"

    def boom(messages):
        raise ap.llm.LLMError("connection refused")

    result = ap.draft_answer(q, r.retrieve(q), chat_fn=boom)
    assert result["Needs Review"] is True
    assert "LLM UNAVAILABLE" in result["Draft Answer"]


# --- Output generation -------------------------------------------------------
def test_run_generates_outputs_with_expected_columns(tmp_path):
    results = ap.run(
        str(Path(ap.__file__).parent / "input" / "questions.csv"),
        POLICIES,
        str(tmp_path),
        chat_fn=fake_chat(),
    )
    assert len(results) == 15

    csv_path = tmp_path / "answers.csv"
    xlsx_path = tmp_path / "answers.xlsx"
    report_path = tmp_path / "review-report.md"
    assert csv_path.exists() and xlsx_path.exists() and report_path.exists()

    with open(csv_path, encoding="utf-8") as f:
        first = f.readline()
        assert "AI-GENERATED" in first  # disclaimer header line
        header = next(csv.reader([f.readline()]))
    assert header == ap.COLUMNS

    from openpyxl import load_workbook

    ws = load_workbook(xlsx_path).active
    assert "AI-GENERATED" in str(ws.cell(1, 1).value)
    assert [c.value for c in ws[2]] == ap.COLUMNS

    assert "AI-GENERATED" in report_path.read_text(encoding="utf-8")


def test_read_questions_supports_xlsx(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["ID", "Question"])
    ws.append(["Q1", "How are encryption keys rotated?"])
    p = tmp_path / "q.xlsx"
    wb.save(p)
    assert ap.read_questions(str(p)) == ["How are encryption keys rotated?"]
